"""Constructor del .docx del documento de arquitectura E4 (entregable 5).

Uso: python construir_docx.py contenido.md salida.docx
Las rutas de imágenes se resuelven relativas al .md.
Especificación de formato: ver README.md de esta carpeta.
"""
import re
import struct
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT  # noqa: F401  (A4 vertical, por claridad)
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FUENTE = "Inter"
ACENTO = "0F766E"
TINTE = "F0FDFA"
TITULO = "0F172A"
TEXTO = "334155"
SECUNDARIO = "475569"
LINEA = "CBD5E1"
ENCABEZADO = "F1F5F9"
AMBAR = "B45309"
AMBAR_FONDO = "FFFBEB"
ANCHO_TEXTO = 15.0  # cm

INTEGRANTES = [
    ("Nicolás Delorte", "172.817-9"),
    ("Camila Belén Lencina", "215.042-6"),
    ("Lara Anush Eriakian", "213.678-8"),
    ("Tadeo Sorrentino", "214.153-0"),
    ("Sofia Maria Deane", "210.350-3"),
    ("Martín Ilán Zajdenberg", "209.501-4"),
    ("Milton Christopher Bernardo Estigarribia", "208.731-5"),
    ("Miranda Rossi", "208.818-6"),
    ("Camila Aylen Suarez", "214.170-0"),
]
VERSION = "Versión 4.0 · 9 de octubre de 2026"
PIE = "Arquitectura del sistema DonaTrack — Entrega 4 · v4.0"


def rgb(h):
    return RGBColor.from_string(h)


# ---------------------------------------------------------------- XML helpers
def el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn("w:" + k), str(v))
    return e


def fijar_fuente(rpr):
    """rFonts explícito en los 4 atributos, sin atributos de tema."""
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = el("w:rFonts")
        rpr.insert(0, rf)
    for a in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        rf.attrib.pop(qn("w:" + a), None)
    for a in ("ascii", "hAnsi", "eastAsia", "cs"):
        rf.set(qn("w:" + a), FUENTE)


def fijar_idioma(rpr):
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = el("w:lang")
        rpr.append(lang)
    lang.set(qn("w:val"), "es-AR")


def campo(par, instr, placeholder=""):
    for tipo in ("begin", "instr", "separate", "texto", "end"):
        r = par.add_run()
        if tipo == "instr":
            t = el("w:instrText")
            t.set(qn("xml:space"), "preserve")
            t.text = f" {instr} "
            r._r.append(t)
        elif tipo == "texto":
            r.text = placeholder
        else:
            r._r.append(el("w:fldChar", fldCharType=tipo))


# ---------------------------------------------------------------- estilos y página
def configurar(doc):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.5)
    sec.left_margin = sec.right_margin = Cm(3.0)
    sec.header_distance = sec.footer_distance = Cm(1.25)

    # docDefaults sin fuentes de tema
    rpr_def = doc.styles.element.find(qn("w:docDefaults") + "/" + qn("w:rPrDefault") + "/" + qn("w:rPr"))
    if rpr_def is not None:
        fijar_fuente(rpr_def)
        fijar_idioma(rpr_def)

    st = doc.styles
    n = st["Normal"]
    n.font.size = Pt(11)
    n.font.color.rgb = rgb(TEXTO)
    pf = n.paragraph_format
    pf.line_spacing = 1.25
    pf.space_after = Pt(6)
    pf.space_before = Pt(0)
    pf.widow_control = True
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT

    for nombre, sz, color, antes, despues in (
        ("Heading 1", 18, TITULO, 18, 6),
        ("Heading 2", 13, ACENTO, 12, 4),
    ):
        s = st[nombre]
        s.font.size = Pt(sz)
        s.font.bold = True
        s.font.italic = False
        s.font.color.rgb = rgb(color)
        p = s.paragraph_format
        p.space_before, p.space_after = Pt(antes), Pt(despues)
        p.keep_with_next = True
        p.widow_control = True
        p.line_spacing = 1.15

    c = st["Caption"]
    c.font.size = Pt(9.5)
    c.font.bold = False
    c.font.italic = False
    c.font.color.rgb = rgb(SECUNDARIO)
    c.paragraph_format.space_before = Pt(4)
    c.paragraph_format.space_after = Pt(10)

    for nombre in ("List Bullet", "List Bullet 2"):
        st[nombre].paragraph_format.space_after = Pt(3)

    for nombre in ("Normal", "Heading 1", "Heading 2", "Caption", "List Bullet", "List Bullet 2"):
        rpr = st[nombre].element.get_or_add_rPr()
        fijar_fuente(rpr)
        fijar_idioma(rpr)

    u = el("w:updateFields", val="true")
    doc.settings.element.append(u)


def pie(doc):
    sec = doc.sections[0]
    sec.different_first_page_header_footer = True
    sec.first_page_footer.is_linked_to_previous = False  # portada sin pie
    p = sec.footer.paragraphs[0]
    p.style = doc.styles["Normal"]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(ANCHO_TEXTO), WD_TAB_ALIGNMENT.RIGHT)
    p.add_run(PIE + "\tPágina ")
    campo(p, "PAGE", "1")
    for r in p.runs:
        r.font.size = Pt(9)
        r.font.color.rgb = rgb(SECUNDARIO)


# ---------------------------------------------------------------- inline
INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*\s][^*]*\*)")


def inline(par, texto, size=None, color=None, bold=False):
    for trozo in INLINE.split(texto):
        if not trozo:
            continue
        if trozo.startswith("**") and trozo.endswith("**"):
            r = par.add_run(trozo[2:-2]); r.bold = True
        elif trozo.startswith("*") and trozo.endswith("*") and len(trozo) > 2:
            r = par.add_run(trozo[1:-1]); r.italic = True
            if bold:
                r.bold = True
        else:
            r = par.add_run(trozo)
            if bold:
                r.bold = True
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = rgb(color)
    return par


# ---------------------------------------------------------------- tablas
def _tblpr_insertar(tblpr, nuevo, antes_de):
    viejo = tblpr.find(nuevo.tag)
    if viejo is not None:
        tblpr.remove(viejo)
    for tag in antes_de:
        ref = tblpr.find(qn(tag))
        if ref is not None:
            ref.addprevious(nuevo)
            return
    tblpr.append(nuevo)


def preparar_tabla(tabla, anchos, bordes):
    """anchos en cm; bordes: dict lado -> (sz, color) o None."""
    tabla.autofit = False
    tabla.alignment = WD_TABLE_ALIGNMENT.LEFT
    tblpr = tabla._tbl.tblPr
    total = int(sum(anchos) * 567)
    tblw = tblpr.find(qn("w:tblW"))
    tblw.set(qn("w:type"), "dxa"); tblw.set(qn("w:w"), str(total))
    tb = el("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        v = bordes.get(lado)
        if v:
            tb.append(el("w:" + lado, val="single", sz=v[0], space=0, color=v[1]))
        else:
            tb.append(el("w:" + lado, val="nil"))
    _tblpr_insertar(tblpr, tb, ("w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook"))
    mar = el("w:tblCellMar")
    for lado, v in (("top", 57), ("left", 85), ("bottom", 57), ("right", 85)):
        mar.append(el("w:" + lado, w=v, type="dxa"))
    _tblpr_insertar(tblpr, mar, ("w:tblLook", "w:tblCaption"))
    for i, col in enumerate(tabla.columns):
        col.width = Cm(anchos[i])
    for fila in tabla.rows:
        for i, celda in enumerate(fila.cells):
            celda.width = Cm(anchos[min(i, len(anchos) - 1)])


def fila_props(fila, encabezado=False):
    trpr = fila._tr.get_or_add_trPr()
    trpr.append(el("w:cantSplit"))
    if encabezado:
        trpr.append(el("w:tblHeader"))


def celda_formato(celda, fondo=None, borde_izq=None):
    """Orden de esquema tcPr: tcW → (gridSpan) → tcBorders → shd → vAlign."""
    tcpr = celda._tc.get_or_add_tcPr()
    for tag in ("w:tcBorders", "w:shd"):
        v = tcpr.find(qn(tag))
        if v is not None:
            tcpr.remove(v)
    valign = tcpr.find(qn("w:vAlign"))
    nuevos = []
    if borde_izq:
        b = el("w:tcBorders")
        b.append(el("w:left", val="single", sz=borde_izq[0], space=0, color=borde_izq[1]))
        nuevos.append(b)
    if fondo:
        nuevos.append(el("w:shd", val="clear", color="auto", fill=fondo))
    for n in nuevos:
        if valign is not None:
            valign.addprevious(n)
        else:
            tcpr.append(n)


def parrafo_celda(celda, primero=True):
    p = celda.paragraphs[0] if primero and not celda.paragraphs[0].text and not celda.paragraphs[0].runs else celda.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    return p


def epigrafe(doc, rotulo, texto, keep_next):
    p = doc.add_paragraph(style="Caption")
    r = p.add_run(rotulo)
    r.bold = True
    r.font.color.rgb = rgb(ACENTO)
    if texto:
        inline(p, " " + texto)
    p.paragraph_format.keep_with_next = keep_next
    if keep_next:
        p.paragraph_format.space_after = Pt(4)
    return p


def separar_celdas(linea):
    s = linea.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|") and not s.endswith("\\|"):
        s = s[:-1]
    partes = re.split(r"(?<!\\)\|", s)
    return [p.strip().replace("\\|", "|") for p in partes]


def tabla_gfm(doc, filas_md, anchos):
    filas = [separar_celdas(l) for l in filas_md if not re.match(r"^\s*\|?\s*:?-{2,}", l)]
    ncols = max(len(f) for f in filas)
    if not anchos or len(anchos) != ncols:
        anchos = [ANCHO_TEXTO / ncols] * ncols
    t = doc.add_table(rows=len(filas), cols=ncols)
    preparar_tabla(t, anchos, {"top": (4, LINEA), "bottom": (4, LINEA), "insideH": (4, LINEA)})
    for i, fila in enumerate(filas):
        row = t.rows[i]
        fila_props(row, encabezado=(i == 0))
        for j in range(ncols):
            c = row.cells[j]
            texto = fila[j] if j < len(fila) else ""
            p = parrafo_celda(c)
            if i == 0:
                inline(p, texto, size=10, color=TITULO, bold=True)
                celda_formato(c, fondo=ENCABEZADO)
            else:
                inline(p, texto, size=10)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


# ---------------------------------------------------------------- recuadros
ROTULO = re.compile(r"^([A-ZÁÉÍÓÚÑ][\wáéíóúñÁÉÍÓÚÑ ]{0,24}):\s*(.*)$")


def recuadro_decision(doc, cabecera, cuerpo):
    partes = [p.strip() for p in cabecera.split("|")]
    ident = partes[0] if partes else ""
    titulo = partes[1] if len(partes) > 1 else ""
    estado = partes[2] if len(partes) > 2 else ""
    bloques = []  # (rotulo, [(tipo, nivel, texto)])
    for linea in cuerpo:
        if not linea.strip():
            continue
        m_b = re.match(r"^(\s*)[-*]\s+(.*)$", linea)
        m_r = ROTULO.match(linea.strip()) if not m_b else None
        if m_r:
            bloques.append((m_r.group(1), []))
            if m_r.group(2):
                bloques[-1][1].append(("p", 0, m_r.group(2)))
        elif m_b:
            if not bloques:
                bloques.append(("", []))
            bloques[-1][1].append(("li", 1 if len(m_b.group(1)) >= 2 else 0, m_b.group(2)))
        else:
            if not bloques:
                bloques.append(("", []))
            bloques[-1][1].append(("p", 0, linea.strip()))

    t = doc.add_table(rows=1 + len(bloques), cols=2)
    anchos = [3.2, 11.8]
    preparar_tabla(t, anchos, {})
    titulo_celda = t.cell(0, 0).merge(t.cell(0, 1))
    pars = []
    p = parrafo_celda(titulo_celda)
    # sin identificador ni rótulo de estado en el título: el estado se cuenta en el texto
    r = p.add_run(f"Decisión: {titulo}")
    r.bold = True; r.font.size = Pt(11); r.font.color.rgb = rgb(ACENTO)
    pars.append(p)
    for i, (rot, items) in enumerate(bloques, start=1):
        c0, c1 = t.rows[i].cells[0], t.rows[i].cells[1]
        p = parrafo_celda(c0)
        r = p.add_run(rot); r.bold = True; r.font.size = Pt(10); r.font.color.rgb = rgb(TITULO)
        pars.append(p)
        primero = True
        for tipo, nivel, texto in items:
            p = parrafo_celda(c1, primero)
            primero = False
            if tipo == "li":
                p.style = doc.styles["List Bullet 2" if nivel else "List Bullet"]
                p.paragraph_format.space_after = Pt(1)
            inline(p, texto, size=10)
            pars.append(p)
    for fila in t.rows:
        fila_props(fila)
    for fila in t.rows:
        vistas = set()
        for j, c in enumerate(fila.cells):
            if id(c._tc) in vistas:
                continue
            vistas.add(id(c._tc))
            celda_formato(c, fondo=TINTE, borde_izq=(24, ACENTO) if j == 0 else None)
    for p in pars[:-1]:
        p.paragraph_format.keep_with_next = True
    for fila in t.rows:
        for c in fila.cells:
            for p in c.paragraphs:
                if p is not pars[-1] and p._p is not pars[-1]._p:
                    p.paragraph_format.keep_with_next = True
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def llamada(doc, rotulo, cuerpo):
    t = doc.add_table(rows=1, cols=1)
    preparar_tabla(t, [ANCHO_TEXTO], {})
    fila_props(t.rows[0])
    c = t.cell(0, 0)
    celda_formato(c, fondo=AMBAR_FONDO, borde_izq=(24, AMBAR))
    texto = " ".join(l.strip() for l in cuerpo if l.strip())
    p = parrafo_celda(c)
    r = p.add_run(rotulo.strip().upper() + ". ")
    r.bold = True; r.font.size = Pt(10); r.font.color.rgb = rgb(AMBAR)
    inline(p, texto, size=10)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


# ---------------------------------------------------------------- figuras
def ancho_png(ruta):
    with open(ruta, "rb") as f:
        cab = f.read(24)
    if cab[:8] != b"\x89PNG\r\n\x1a\n":
        return None, None
    return struct.unpack(">II", cab[16:24])


def figura(doc, ruta, alt):
    w_px, h_px = ancho_png(ruta)
    ancho = ANCHO_TEXTO if not w_px else min(ANCHO_TEXTO, w_px * ANCHO_TEXTO / 730)
    if w_px and h_px and ancho * h_px / w_px > 20.0:
        ancho = 20.0 * w_px / h_px
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    p.add_run().add_picture(str(ruta), width=Cm(ancho))
    m = re.match(r"^(Figura\s+\d+\.)\s*(.*)$", alt)
    if m:
        cap = epigrafe(doc, m.group(1), m.group(2), keep_next=False)
    else:
        cap = doc.add_paragraph(style="Caption"); inline(cap, alt)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER


# ---------------------------------------------------------------- portada e índice
def portada(doc):
    for _ in range(3):
        doc.add_paragraph()

    def linea(txt, sz, bold=False, color=None, antes=0, despues=6):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(antes)
        p.paragraph_format.space_after = Pt(despues)
        r = p.add_run(txt)
        r.font.size = Pt(sz); r.bold = bold
        if color:
            r.font.color.rgb = rgb(color)
        return p

    linea("Universidad Tecnológica Nacional — Facultad Regional Buenos Aires", 12)
    linea("Diseño de Sistemas · Curso K3002 · 2026", 12, despues=36)
    linea("Arquitectura del sistema DonaTrack", 28, bold=True, color=TITULO, despues=10).paragraph_format.line_spacing = 1.0
    linea("Entrega 4 — Persistencia, integración y despliegue", 14, color=SECUNDARIO)
    linea("Documento explicativo de la arquitectura (entregable 5)", 11, despues=30)
    linea("Grupo 5", 14, bold=True, color=TITULO, despues=10)

    t = doc.add_table(rows=1 + len(INTEGRANTES), cols=2)
    preparar_tabla(t, [10.5, 4.5], {"top": (4, LINEA), "bottom": (4, LINEA), "insideH": (4, LINEA)})
    for i, (nombre, legajo) in enumerate([("Nombre", "Legajo")] + INTEGRANTES):
        fila = t.rows[i]
        fila_props(fila, encabezado=(i == 0))
        for j, txt in enumerate((nombre, legajo)):
            c = fila.cells[j]
            p = parrafo_celda(c)
            r = p.add_run(txt); r.font.size = Pt(10)
            if i == 0:
                r.bold = True; r.font.color.rgb = rgb(TITULO)
                celda_formato(c, fondo=ENCABEZADO)
    linea(VERSION, 11, color=SECUNDARIO, antes=48)

    p = doc.add_paragraph()
    p.paragraph_format.page_break_before = True
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run("Índice")
    r.bold = True; r.font.size = Pt(18); r.font.color.rgb = rgb(TITULO)
    p = doc.add_paragraph()
    campo(p, r'TOC \o "1-2" \h \z \u', "Actualizar el índice: clic derecho → Actualizar campo (en Google Docs: Insertar → Índice).")


# ---------------------------------------------------------------- parser principal
def salto_antes(nivel, texto, primer_capitulo):
    if nivel == 1:
        return primer_capitulo or texto.startswith("Anexo") or re.match(r"^11\b", texto) is not None
    # 11.1 comparte página con la intro del capítulo 11; las demás fichas empiezan en página nueva
    return re.match(r"^11\.([2-9]|\d{2,})\b", texto) is not None


def construir(src, dst):
    base = src.parent
    crudo = src.read_text(encoding="utf-8")
    crudo = re.sub(r"<!--.*?-->", "", crudo, flags=re.S)
    lineas = crudo.splitlines()

    doc = Document()
    configurar(doc)
    pie(doc)
    portada(doc)

    ultimo = {"p": None}  # último párrafo de texto (para keep_with_next de la oración introductoria)
    primer_capitulo = True
    i = 0
    buffer = []

    def volcar():
        if buffer:
            p = doc.add_paragraph()
            inline(p, " ".join(s.strip() for s in buffer))
            ultimo["p"] = p
            buffer.clear()

    def atar_intro():
        if ultimo["p"] is not None:
            ultimo["p"].paragraph_format.keep_with_next = True
        ultimo["p"] = None

    while i < len(lineas):
        linea = lineas[i]
        s = linea.strip()

        if not s:
            volcar(); i += 1; continue

        m = re.match(r"^(#{1,2})\s+(.*)$", s)
        if m:
            volcar()
            nivel = len(m.group(1))
            texto = m.group(2).strip()
            h = doc.add_heading(level=nivel)
            inline(h, texto)
            if salto_antes(nivel, texto, primer_capitulo and nivel == 1):
                h.paragraph_format.page_break_before = True
            if nivel == 1:
                primer_capitulo = False
            ultimo["p"] = None
            i += 1; continue

        m = re.match(r"^:::(decision|decisión|llamada)\s*(.*)$", s, re.I)
        if m:
            volcar(); atar_intro()
            cuerpo = []
            i += 1
            while i < len(lineas) and lineas[i].strip() != ":::":
                cuerpo.append(lineas[i]); i += 1
            i += 1
            if m.group(1).lower().startswith("decis"):
                recuadro_decision(doc, m.group(2), cuerpo)
            else:
                llamada(doc, m.group(2), cuerpo)
            continue

        m = re.match(r"^!\[(.*?)\]\((.*?)\)\s*$", s)
        if m:
            volcar(); atar_intro()
            figura(doc, (base / m.group(2)).resolve(), m.group(1).strip())
            i += 1; continue

        m = re.match(r"^(Tabla\s+\d+\.)\s*(.*)$", s)
        if m and not buffer:
            j = i + 1
            while j < len(lineas) and not lineas[j].strip():
                j += 1
            anchos = None
            ma = re.match(r"^\{\s*anchos\s*:\s*(.*)\}\s*$", lineas[j].strip()) if j < len(lineas) else None
            if ma:
                anchos = [float(x) for x in re.findall(r"\d+(?:\.\d+)?", ma.group(1))]
                j += 1
                while j < len(lineas) and not lineas[j].strip():
                    j += 1
            if j < len(lineas) and lineas[j].strip().startswith("|"):
                atar_intro()
                epigrafe(doc, m.group(1), m.group(2), keep_next=True)
                filas = []
                while j < len(lineas) and lineas[j].strip().startswith("|"):
                    filas.append(lineas[j]); j += 1
                tabla_gfm(doc, filas, anchos)
                i = j; continue

        if s.startswith("|"):
            volcar(); atar_intro()
            filas = []
            while i < len(lineas) and lineas[i].strip().startswith("|"):
                filas.append(lineas[i]); i += 1
            tabla_gfm(doc, filas, None)
            continue

        m = re.match(r"^(\s*)[-*]\s+(.*)$", linea)
        if m:
            volcar(); atar_intro()
            nivel2 = len(m.group(1).replace("\t", "  ")) >= 2
            p = doc.add_paragraph(style="List Bullet 2" if nivel2 else "List Bullet")
            inline(p, m.group(2).strip())
            i += 1; continue

        if re.match(r"^\{\s*anchos\s*:", s):  # anchos huérfanos
            i += 1; continue

        buffer.append(s)
        i += 1
    volcar()
    doc.save(dst)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit("Uso: python construir_docx.py contenido.md salida.docx")
    construir(Path(sys.argv[1]), Path(sys.argv[2]))
    print(f"OK: {sys.argv[2]}")
